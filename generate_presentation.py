import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Constants
    BG_DARK = RGBColor(11, 15, 23)        # #0B0F17 - Deep Navy Canvas
    CARD_BG = RGBColor(22, 30, 46)        # #161E2E - Standard Card Surface
    CARD_BG_ALT = RGBColor(17, 24, 39)    # #111827 - Secondary Card Surface
    CARD_BORDER = RGBColor(46, 61, 86)    # #2E3D56 - Border

    CYAN_PRIMARY = RGBColor(56, 189, 248) # #38BDF8 - Vibrant Cyber Cyan
    BLUE_BADGE = RGBColor(23, 42, 70)     # #172A46 - Badge Surface

    SAFE_GREEN = RGBColor(16, 185, 129)   # #10B981 - Safe 🟢
    WARN_AMBER = RGBColor(245, 158, 11)   # #F59E0B - Suspicious 🟠
    DANGER_RED = RGBColor(239, 68, 68)    # #EF4444 - High Risk 🔴

    TEXT_TITLE = RGBColor(248, 250, 252)  # #F8FAFC - Pure White
    TEXT_BODY = RGBColor(203, 213, 225)   # #CBD5E1 - Light Slate Body
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8 - Subtitle / Muted Slate
    TEXT_DIM = RGBColor(100, 116, 139)    # #64748B - Subtle Footers

    FONT_MAIN = "Arial"

    # Exact Assets
    LOGO_BADGE = "presentation_assets/logo_app_badge.png"
    LOGO_WHITE = "presentation_assets/logo_white_trans.png"

    PHONE_DASHBOARD = "presentation_assets/mockups/phone_dashboard_exact.png"
    PHONE_THREAT = "presentation_assets/mockups/phone_threat_detail_exact.png"
    PHONE_LANGUAGES = "presentation_assets/mockups/phone_languages_exact.png"
    PHONE_SPLASH = "presentation_assets/mockups/phone_splash_exact.png"

    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text=None):
        # Badge Pill
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(3.4), Inches(0.32)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = BLUE_BADGE
        badge.line.color.rgb = CYAN_PRIMARY
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        r_b = p_b.add_run()
        r_b.text = badge_text.upper()
        r_b.font.name = FONT_MAIN
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = CYAN_PRIMARY

        # Header Textbox
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(11.733), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = title_text
        r1.font.name = FONT_MAIN
        r1.font.size = Pt(25)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_TITLE

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(3)
            r2 = p2.add_run()
            r2.text = subtitle_text
            r2.font.name = FONT_MAIN
            r2.font.size = Pt(12)
            r2.font.color.rgb = TEXT_MUTED

    def add_footer(slide, slide_num, total_slides=12):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Pt(1))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

        tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(8.0), Inches(0.28))
        tf_l = tb_l.text_frame
        tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
        p_l = tf_l.paragraphs[0]
        r_l = p_l.add_run()
        r_l.text = "SCAMSHIELD 🛡️  •  Hackatopia 2K26 National-Level Hackathon Final Presentation"
        r_l.font.name = FONT_MAIN
        r_l.font.size = Pt(9)
        r_l.font.color.rgb = TEXT_DIM

        tb_r = slide.shapes.add_textbox(Inches(10.5), Inches(7.12), Inches(2.033), Inches(0.28))
        tf_r = tb_r.text_frame
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
        p_r = tf_r.paragraphs[0]
        p_r.alignment = PP_ALIGN.RIGHT
        r_r = p_r.add_run()
        r_r.text = f"{slide_num:02d} / {total_slides:02d}"
        r_r.font.name = FONT_MAIN
        r_r.font.size = Pt(9.5)
        r_r.font.bold = True
        r_r.font.color.rgb = CYAN_PRIMARY

    def make_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
        return card

    # ==========================================
    # SLIDE 1 — TITLE / HERO SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_bg(s1)

    # Hackathon Tag Badge
    tag1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.75), Inches(4.6), Inches(0.36))
    tag1.fill.solid()
    tag1.fill.fore_color.rgb = BLUE_BADGE
    tag1.line.color.rgb = CYAN_PRIMARY
    tag1.line.width = Pt(1)
    tf1 = tag1.text_frame
    p1 = tf1.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "HACKATOPIA 2K26 • NATIONAL-LEVEL HACKATHON FINALS"
    r1.font.name = FONT_MAIN
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = CYAN_PRIMARY

    # App Logo Badge (High-res crisp icon)
    if os.path.exists(LOGO_BADGE):
        s1.shapes.add_picture(LOGO_BADGE, Inches(0.8), Inches(1.35), width=Inches(1.25), height=Inches(1.25))

    # Main Titles
    tb_title = s1.shapes.add_textbox(Inches(2.25), Inches(1.32), Inches(5.8), Inches(1.3))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = "SCAMSHIELD"
    r_t.font.name = FONT_MAIN
    r_t.font.size = Pt(44)
    r_t.font.bold = True
    r_t.font.color.rgb = TEXT_TITLE

    p_tag = tf_t.add_paragraph()
    p_tag.space_before = Pt(3)
    r_tag = p_tag.add_run()
    r_tag.text = "Detect. Explain. Protect."
    r_tag.font.name = FONT_MAIN
    r_tag.font.size = Pt(16)
    r_tag.font.bold = True
    r_tag.font.color.rgb = CYAN_PRIMARY

    # Subtitle Paragraph Box
    tb_sub = s1.shapes.add_textbox(Inches(0.8), Inches(2.78), Inches(6.8), Inches(1.4))
    tf_s = tb_sub.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    r_s = p_s.add_run()
    r_s.text = "AI-Powered Real-Time Protection Against Mobile Scam Messages"
    r_s.font.name = FONT_MAIN
    r_s.font.size = Pt(20)
    r_s.font.bold = True
    r_s.font.color.rgb = TEXT_TITLE

    p_desc = tf_s.add_paragraph()
    p_desc.space_before = Pt(8)
    r_desc = p_desc.add_run()
    r_desc.text = "A zero-touch cybersecurity companion that automatically intercepts incoming SMS & WhatsApp threats on physical Android devices — delivering instant explainable risk analysis before victims can be deceived."
    r_desc.font.name = FONT_MAIN
    r_desc.font.size = Pt(13)
    r_desc.font.color.rgb = TEXT_BODY

    # 3 Value Highlights
    pills = [
        ("🛡️ Zero-Copy UX", "Automatic background interception on device"),
        ("🧠 Dual-Engine AI", "Edge heuristics + Groq LLaMA 3.3 70B cloud"),
        ("👴 Elderly-First", "High contrast, trilingual (EN/KN/HI) & audio TTS")
    ]
    for i, (p_title, p_sub) in enumerate(pills):
        top_y = Inches(4.45 + i * 0.72)
        make_card(s1, Inches(0.8), top_y, Inches(6.8), Inches(0.62), CARD_BG, CARD_BORDER)
        tb_p = s1.shapes.add_textbox(Inches(1.0), top_y + Inches(0.06), Inches(6.4), Inches(0.5))
        tf_p = tb_p.text_frame
        p_pt = tf_p.paragraphs[0]
        r_pt = p_pt.add_run()
        r_pt.text = p_title + "  —  "
        r_pt.font.name = FONT_MAIN
        r_pt.font.size = Pt(12)
        r_pt.font.bold = True
        r_pt.font.color.rgb = CYAN_PRIMARY
        
        r_ps = p_pt.add_run()
        r_ps.text = p_sub
        r_ps.font.name = FONT_MAIN
        r_ps.font.size = Pt(11.5)
        r_ps.font.color.rgb = TEXT_BODY

    # Team Box
    make_card(s1, Inches(0.8), Inches(6.68), Inches(6.8), Inches(0.52), CARD_BG_ALT, CARD_BORDER)
    tb_team = s1.shapes.add_textbox(Inches(1.0), Inches(6.72), Inches(6.4), Inches(0.42))
    tf_team = tb_team.text_frame
    p_tm = tf_team.paragraphs[0]
    r_tm = p_tm.add_run()
    r_tm.text = "Project Team: "
    r_tm.font.name = FONT_MAIN
    r_tm.font.size = Pt(11)
    r_tm.font.bold = True
    r_tm.font.color.rgb = TEXT_MUTED

    r_tm2 = p_tm.add_run()
    r_tm2.text = "Aaron Kuriyan (Lead Developer)  •  @zer0neo (Collaborator)"
    r_tm2.font.name = FONT_MAIN
    r_tm2.font.size = Pt(11)
    r_tm2.font.color.rgb = TEXT_TITLE

    # Right Phone Mockup
    if os.path.exists(PHONE_DASHBOARD):
        s1.shapes.add_picture(PHONE_DASHBOARD, Inches(8.4), Inches(0.75), height=Inches(6.4))

    # ==========================================
    # SLIDE 2 — THE PROBLEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_bg(s2)
    add_header(s2, "01 | Problem Understanding", "One Message Can Cost Someone Everything.", 
               "Digital scams exploit urgency, authority, and fear — striking directly on the apps users trust most.")
    add_footer(s2, 2)

    scam_cards = [
        ("🔴 Fake KYC Deactivation Threat", "SMS • Banking Impersonation", 
         "\"Your bank account will be blocked today. Verify your KYC immediately using this link: http://sbi-kyc-update.net\"", DANGER_RED),
        ("🔴 UPI Collect Request Scam", "WhatsApp • Financial Fraud", 
         "\"Approve collect request of ₹5,000 on PhonePe to receive your government cashback reward now.\"", DANGER_RED),
        ("🟠 Fake Courier & Delivery Fee", "SMS • Phishing Link", 
         "\"Your parcel delivery failed. Pay ₹50 reschedule fee using this link: http://track-parcel-redeliver.in\"", WARN_AMBER),
        ("🔴 Electricity Power Cut Extortion", "SMS • Utility Impersonation", 
         "\"Dear Consumer, your electricity power will be disconnected tonight at 9:30 PM due to unpaid bill. Call officer now.\"", DANGER_RED)
    ]

    for i, (title, sub, msg, bord_c) in enumerate(scam_cards):
        top_y = Inches(1.95 + i * 1.2)
        make_card(s2, Inches(0.8), top_y, Inches(6.8), Inches(1.1), CARD_BG, bord_c, 1.5)
        
        tb = s2.shapes.add_textbox(Inches(1.0), top_y + Inches(0.08), Inches(6.4), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title + "  "
        r_t.font.name = FONT_MAIN
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = bord_c

        r_sub = p_t.add_run()
        r_sub.text = f"({sub})"
        r_sub.font.name = FONT_MAIN
        r_sub.font.size = Pt(10)
        r_sub.font.color.rgb = TEXT_MUTED

        p_m = tf.add_paragraph()
        p_m.space_before = Pt(4)
        r_m = p_m.add_run()
        r_m.text = msg
        r_m.font.name = FONT_MAIN
        r_m.font.size = Pt(11)
        r_m.font.italic = True
        r_m.font.color.rgb = TEXT_BODY

    # Right Column: The Core Vulnerability Analysis Card
    make_card(s2, Inches(8.0), Inches(1.95), Inches(4.533), Inches(4.7), CARD_BG, CARD_BORDER)

    tb_r2 = s2.shapes.add_textbox(Inches(8.25), Inches(2.1), Inches(4.05), Inches(3.2))
    tf_r2 = tb_r2.text_frame
    tf_r2.word_wrap = True

    p_rh = tf_r2.paragraphs[0]
    r_rh = p_rh.add_run()
    r_rh.text = "Why The Crisis Is Escalating"
    r_rh.font.name = FONT_MAIN
    r_rh.font.size = Pt(17)
    r_rh.font.bold = True
    r_rh.font.color.rgb = CYAN_PRIMARY

    bullets_s2 = [
        ("Exploitation of Psychology", "Scammers bypass firewalls by manipulating emotions. Artificial urgency forces instant panic."),
        ("Elderly Citizens Hit Hardest", "Over 80% of elderly mobile users cannot spot spoofed URLs or deceptive domain extensions."),
        ("WhatsApp & SMS Dominance", "Threats arrive inside everyday personal channels alongside genuine messages from family."),
        ("Massive Financial Damage", "Over [INSERT VERIFIED STATISTIC] in personal savings wiped out annually across mobile banking victims.")
    ]

    for b_title, b_desc in bullets_s2:
        p_b = tf_r2.add_paragraph()
        p_b.space_before = Pt(8)
        r_bt = p_b.add_run()
        r_bt.text = f"• {b_title}: "
        r_bt.font.name = FONT_MAIN
        r_bt.font.size = Pt(11)
        r_bt.font.bold = True
        r_bt.font.color.rgb = TEXT_TITLE

        r_bd = p_b.add_run()
        r_bd.text = b_desc
        r_bd.font.name = FONT_MAIN
        r_bd.font.size = Pt(10.5)
        r_bd.font.color.rgb = TEXT_BODY

    # Bottom Callout Box cleanly separated inside card
    stat_box = make_card(s2, Inches(8.2), Inches(5.65), Inches(4.133), Inches(0.85), CARD_BG_ALT, CYAN_PRIMARY, 1)
    tb_st = s2.shapes.add_textbox(Inches(8.35), Inches(5.72), Inches(3.85), Inches(0.72))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True
    p_st = tf_st.paragraphs[0]
    r_st1 = p_st.add_run()
    r_st1.text = "KEY TAKEAWAY: "
    r_st1.font.name = FONT_MAIN
    r_st1.font.size = Pt(10.5)
    r_st1.font.bold = True
    r_st1.font.color.rgb = WARN_AMBER

    r_st2 = p_st.add_run()
    r_st2.text = "Scam messages look identical to official notices. Traditional security expects victims to be cybersecurity experts."
    r_st2.font.name = FONT_MAIN
    r_st2.font.size = Pt(10)
    r_st2.font.color.rgb = TEXT_TITLE

    # ==========================================
    # SLIDE 3 — EXISTING SOLUTIONS & THE GAP
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_bg(s3)
    add_header(s3, "02 | Existing Solutions & Problem Gap", "Today's Protection is Too Reactive.",
               "Existing security applications fail the very users who need protection the most.")
    add_footer(s3, 3)

    # Left: Traditional Friction Workflow
    make_card(s3, Inches(0.8), Inches(1.95), Inches(5.5), Inches(3.95), CARD_BG, CARD_BORDER)
    tb_flow = s3.shapes.add_textbox(Inches(1.05), Inches(2.1), Inches(5.0), Inches(3.65))
    tf_flow = tb_flow.text_frame
    tf_flow.word_wrap = True

    p_fh = tf_flow.paragraphs[0]
    r_fh = p_fh.add_run()
    r_fh.text = "The Traditional 6-Step Copy-Paste Burden"
    r_fh.font.name = FONT_MAIN
    r_fh.font.size = Pt(16)
    r_fh.font.bold = True
    r_fh.font.color.rgb = WARN_AMBER

    steps = [
        ("1. Message Arrives", "User receives deceptive SMS/WhatsApp threat"),
        ("2. User Must Suspect", "Requires prior technical awareness and doubt"),
        ("3. Manually Copy Text", "Long-press message, select text, find copy icon"),
        ("4. Switch Application", "Exit WhatsApp, find security app, open it"),
        ("5. Paste & Submit", "Locate text input, paste message, tap scan"),
        ("6. Decipher Jargon", "Read technical score without actionable guidance")
    ]
    for s_num, s_desc in steps:
        p_s = tf_flow.add_paragraph()
        p_s.space_before = Pt(6)
        r_sn = p_s.add_run()
        r_sn.text = f"{s_num}  ➔  "
        r_sn.font.name = FONT_MAIN
        r_sn.font.size = Pt(11.5)
        r_sn.font.bold = True
        r_sn.font.color.rgb = TEXT_TITLE

        r_sd = p_s.add_run()
        r_sd.text = s_desc
        r_sd.font.name = FONT_MAIN
        r_sd.font.size = Pt(10.5)
        r_sd.font.color.rgb = TEXT_MUTED

    # Right: The 4 Fatal Gaps
    gaps = [
        ("❌ Manual & High-Friction", "Requires 6+ deliberate steps across multiple screens. Elderly and panic-stricken users simply never complete it."),
        ("❌ Protection Starts After Suspicion", "If the victim trusts the impersonated bank or government notice, traditional tools are never even opened."),
        ("❌ Cryptic Threat Jargon", "Displaying 'Phishing Confidence: 87.4%' confuses users instead of giving clear, plain-language directions."),
        ("❌ Invasive Privacy Trade-offs", "Traditional caller-ID and spam tools harvest full address books and upload personal conversation logs to remote clouds.")
    ]
    for i, (g_title, g_desc) in enumerate(gaps):
        top_g = Inches(1.95 + i * 1.0)
        make_card(s3, Inches(6.6), top_g, Inches(5.933), Inches(0.9), CARD_BG, DANGER_RED, 1)
        tb_g = s3.shapes.add_textbox(Inches(6.8), top_g + Inches(0.08), Inches(5.5), Inches(0.75))
        tf_g = tb_g.text_frame
        tf_g.word_wrap = True
        p_gt = tf_g.paragraphs[0]
        r_gt = p_gt.add_run()
        r_gt.text = g_title
        r_gt.font.name = FONT_MAIN
        r_gt.font.size = Pt(12)
        r_gt.font.bold = True
        r_gt.font.color.rgb = DANGER_RED

        p_gd = tf_g.add_paragraph()
        p_gd.space_before = Pt(2)
        r_gd = p_gd.add_run()
        r_gd.text = g_desc
        r_gd.font.name = FONT_MAIN
        r_gd.font.size = Pt(10.5)
        r_gd.font.color.rgb = TEXT_BODY

    # Bottom Banner Transition
    make_card(s3, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.75), BLUE_BADGE, CYAN_PRIMARY, 1.5)
    tb_b3 = s3.shapes.add_textbox(Inches(1.0), Inches(6.16), Inches(11.333), Inches(0.6))
    tf_b3 = tb_b3.text_frame
    tf_b3.word_wrap = True
    p_b3 = tf_b3.paragraphs[0]
    p_b3.alignment = PP_ALIGN.CENTER
    r_b3_1 = p_b3.add_run()
    r_b3_1.text = "THE FATAL PARADOX:  "
    r_b3_1.font.name = FONT_MAIN
    r_b3_1.font.size = Pt(12)
    r_b3_1.font.bold = True
    r_b3_1.font.color.rgb = WARN_AMBER

    r_b3_2 = p_b3.add_run()
    r_b3_2.text = "Traditional tools require users to recognize the scam BEFORE they can be protected. "
    r_b3_2.font.name = FONT_MAIN
    r_b3_2.font.size = Pt(12)
    r_b3_2.font.bold = True
    r_b3_2.font.color.rgb = TEXT_TITLE

    r_b3_3 = p_b3.add_run()
    r_b3_3.text = "What if protection happened automatically the instant a message arrives?"
    r_b3_3.font.name = FONT_MAIN
    r_b3_3.font.size = Pt(12)
    r_b3_3.font.bold = True
    r_b3_3.font.color.rgb = CYAN_PRIMARY

    # ==========================================
    # SLIDE 4 — OUR PROPOSED SOLUTION
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_bg(s4)
    add_header(s4, "03 | Proposed Solution", "Meet SCAMSHIELD.", 
               "Turning message detection into proactive, zero-touch mobile protection.")
    add_footer(s4, 4)

    pillars = [
        ("🛡️ 1. AUTOMATIC DETECTION", 
         "Zero Copy-Paste Required",
         "Native Android background service captures incoming SMS and WhatsApp messages directly via NotificationListenerService and Telephony receiver. No manual scanning, screenshots, or pasting needed.",
         CYAN_PRIMARY),
        ("🧠 2. INTELLIGENT DUAL-ENGINE AI", 
         "Edge Filtering + Cloud Reasoning",
         "On-device weighted heuristics filter normal chats instantly (<5ms). Suspicious signals undergo deep Groq LLaMA 3.3 70B contextual analysis to identify deceptive intent, urgency, and fraud tactics.",
         SAFE_GREEN),
        ("👴 3. ELDERLY-FIRST ACCESSIBILITY", 
         "Explainable, Trilingual & Voice-Enabled",
         "High-contrast UI designed for non-technical users. Delivers plain-language guidance ('Do NOT click link', 'Do NOT share OTP'), full Kannada and Hindi localization, and Spoken Voice Warnings (TTS).",
         WARN_AMBER)
    ]

    for i, (p_head, p_sub, p_body, col) in enumerate(pillars):
        top_p = Inches(1.95 + i * 1.55)
        make_card(s4, Inches(0.8), top_p, Inches(7.5), Inches(1.4), CARD_BG, col, 1.5)
        
        tb = s4.shapes.add_textbox(Inches(1.05), top_p + Inches(0.1), Inches(7.0), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_h = tf.paragraphs[0]
        r_h = p_h.add_run()
        r_h.text = p_head + "  "
        r_h.font.name = FONT_MAIN
        r_h.font.size = Pt(14)
        r_h.font.bold = True
        r_h.font.color.rgb = col

        r_sub = p_h.add_run()
        r_sub.text = f"— {p_sub}"
        r_sub.font.name = FONT_MAIN
        r_sub.font.size = Pt(11)
        r_sub.font.color.rgb = TEXT_MUTED

        p_b = tf.add_paragraph()
        p_b.space_before = Pt(4)
        r_b = p_b.add_run()
        r_b.text = p_body
        r_b.font.name = FONT_MAIN
        r_b.font.size = Pt(11.5)
        r_b.font.color.rgb = TEXT_BODY

    # Right Phone Mockup
    if os.path.exists(PHONE_SPLASH):
        s4.shapes.add_picture(PHONE_SPLASH, Inches(8.9), Inches(1.6), height=Inches(5.2))

    # ==========================================
    # SLIDE 5 — HOW SCAMSHIELD WORKS (WORKFLOW)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_bg(s5)
    add_header(s5, "04 | How The Solution Works", "From Message to Protection in Under 500ms.",
               "A continuous, real-time background workflow engineered for speed, privacy, and explainability.")
    add_footer(s5, 5)

    steps_s5 = [
        ("01", "INCOMING MESSAGE", "SMS or WhatsApp message arrives on phone", CYAN_PRIMARY),
        ("02", "OS DETECTION", "Android NotificationListener & SmsReceiver", CYAN_PRIMARY),
        ("03", "EDGE PRE-FILTER", "On-device rules filter 90%+ benign messages", SAFE_GREEN),
        ("04", "AI ANALYSIS", "Groq LLaMA 3.3 70B inspects suspicious intent", WARN_AMBER),
        ("05", "RISK CALIBRATION", "0–100 score + threat category generated", DANGER_RED),
        ("06", "USER PROTECTION", "Heads-Up alert + Spoken Voice Warning (TTS)", CYAN_PRIMARY)
    ]

    for i, (num, title, desc, col) in enumerate(steps_s5):
        left_x = Inches(0.8 + i * 1.98)
        make_card(s5, left_x, Inches(2.0), Inches(1.85), Inches(2.2), CARD_BG, col, 1.5)
        
        tb = s5.shapes.add_textbox(left_x + Inches(0.1), Inches(2.1), Inches(1.65), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_n = tf.paragraphs[0]
        r_n = p_n.add_run()
        r_n.text = num
        r_n.font.name = FONT_MAIN
        r_n.font.size = Pt(18)
        r_n.font.bold = True
        r_n.font.color.rgb = col

        p_t = tf.add_paragraph()
        p_t.space_before = Pt(4)
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.name = FONT_MAIN
        r_t.font.size = Pt(11)
        r_t.font.bold = True
        r_t.font.color.rgb = TEXT_TITLE

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(6)
        r_d = p_d.add_run()
        r_d.text = desc
        r_d.font.name = FONT_MAIN
        r_d.font.size = Pt(10)
        r_d.font.color.rgb = TEXT_MUTED

    tiers = [
        ("🟢 SAFE (0–25%)", "Normal Everyday Chats", 
         "Messages like 'Dinner ready' or 'Call you at 5' trigger zero alarms. Stored silently in local inbox. Never leaves user's device.", SAFE_GREEN),
        ("🟠 SUSPICIOUS (26–69%)", "Unverified Links & Deals", 
         "Dubious delivery fees, unknown lottery claims, or unofficial offers. Highlighted with amber caution card and verification guidance.", WARN_AMBER),
        ("🔴 HIGH RISK (70–100%)", "Imminent Financial Threats", 
         "Account blockage threats, OTP phishing, and UPI PIN collect traps. Triggers high-priority Heads-Up Notification and Spoken Audio Alert.", DANGER_RED)
    ]
    for i, (t_title, t_sub, t_desc, col) in enumerate(tiers):
        left_t = Inches(0.8 + i * 3.96)
        make_card(s5, left_t, Inches(4.5), Inches(3.8), Inches(1.6), CARD_BG, col, 1.5)
        
        tb_t = s5.shapes.add_textbox(left_t + Inches(0.15), Inches(4.6), Inches(3.5), Inches(1.4))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        
        p_tt = tf_t.paragraphs[0]
        r_tt = p_tt.add_run()
        r_tt.text = t_title + "  "
        r_tt.font.name = FONT_MAIN
        r_tt.font.size = Pt(12)
        r_tt.font.bold = True
        r_tt.font.color.rgb = col

        r_ts = p_tt.add_run()
        r_ts.text = f"— {t_sub}"
        r_ts.font.name = FONT_MAIN
        r_ts.font.size = Pt(10.5)
        r_ts.font.color.rgb = TEXT_MUTED

        p_td = tf_t.add_paragraph()
        p_td.space_before = Pt(4)
        r_td = p_td.add_run()
        r_td.text = t_desc
        r_td.font.name = FONT_MAIN
        r_td.font.size = Pt(10.5)
        r_td.font.color.rgb = TEXT_BODY

    make_card(s5, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.55), CARD_BG_ALT, SAFE_GREEN, 1)
    tb_pv = s5.shapes.add_textbox(Inches(1.0), Inches(6.35), Inches(11.333), Inches(0.45))
    tf_pv = tb_pv.text_frame
    p_pv = tf_pv.paragraphs[0]
    p_pv.alignment = PP_ALIGN.CENTER
    r_pv1 = p_pv.add_run()
    r_pv1.text = "🔒 PRIVACY-BY-DEFAULT ARCHITECTURE: "
    r_pv1.font.name = FONT_MAIN
    r_pv1.font.size = Pt(11)
    r_pv1.font.bold = True
    r_pv1.font.color.rgb = SAFE_GREEN

    r_pv2 = p_pv.add_run()
    r_pv2.text = "90%+ of harmless messages are resolved locally on the device edge. Zero plaintext personal chats are stored on the server."
    r_pv2.font.name = FONT_MAIN
    r_pv2.font.size = Pt(11)
    r_pv2.font.color.rgb = TEXT_TITLE

    # ==========================================
    # SLIDE 6 — THE ACTUAL APP (PRODUCTION UI)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_bg(s6)
    add_header(s6, "05 | The Actual App", "Protection That Fits Into Everyday Messaging.",
               "Live production UI screenshots running on a physical Android device.")
    add_footer(s6, 6)

    mockup_data = [
        (PHONE_DASHBOARD, "REAL-TIME DASHBOARD", "Unified Message Protection Inbox with live threat counts & instant risk level chips"),
        (PHONE_THREAT, "DETAILED THREAT ANALYSIS", "100% Risk score, why it was flagged, and verified what to do / not to do steps"),
        (PHONE_LANGUAGES, "TRILINGUAL ACCESSIBILITY", "English, Kannada (ಕನ್ನಡ), and Hindi (हिन्दी) system localization + active monitoring toggle")
    ]

    for i, (img_path, label, desc) in enumerate(mockup_data):
        left_m = Inches(0.8 + i * 3.96)
        if os.path.exists(img_path):
            s6.shapes.add_picture(img_path, left_m + Inches(0.55), Inches(1.85), height=Inches(4.1))
        
        make_card(s6, left_m, Inches(6.05), Inches(3.75), Inches(0.85), CARD_BG, CYAN_PRIMARY, 1)
        tb_m = s6.shapes.add_textbox(left_m + Inches(0.1), Inches(6.1), Inches(3.55), Inches(0.75))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        
        p_ml = tf_m.paragraphs[0]
        r_ml = p_ml.add_run()
        r_ml.text = label
        r_ml.font.name = FONT_MAIN
        r_ml.font.size = Pt(11)
        r_ml.font.bold = True
        r_ml.font.color.rgb = CYAN_PRIMARY

        p_md = tf_m.add_paragraph()
        p_md.space_before = Pt(2)
        r_md = p_md.add_run()
        r_md.text = desc
        r_md.font.name = FONT_MAIN
        r_md.font.size = Pt(9.5)
        r_md.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 7 — DETAILED THREAT ANALYSIS
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_bg(s7)
    add_header(s7, "06 | Detailed Threat Analysis", "Don't Just Say 'SCAM'. Explain WHY.",
               "Transforming threat detection into informed, confident user decisions with zero technical jargon.")
    add_footer(s7, 7)

    # Left: Real Phone Mockup of Threat Detail Screen
    if os.path.exists(PHONE_THREAT):
        s7.shapes.add_picture(PHONE_THREAT, Inches(0.8), Inches(1.7), height=Inches(5.2))

    # Right: Structured Explainability Cards
    # Card 1: Live Case Details
    make_card(s7, Inches(3.8), Inches(1.8), Inches(8.733), Inches(1.5), CARD_BG, DANGER_RED, 2)
    tb_cd = s7.shapes.add_textbox(Inches(4.0), Inches(1.9), Inches(8.3), Inches(1.3))
    tf_cd = tb_cd.text_frame
    tf_cd.word_wrap = True
    
    p_cd1 = tf_cd.paragraphs[0]
    r_cd1 = p_cd1.add_run()
    r_cd1.text = "🔴 HIGH RISK  •  100% CONFIDENCE  "
    r_cd1.font.name = FONT_MAIN
    r_cd1.font.size = Pt(14)
    r_cd1.font.bold = True
    r_cd1.font.color.rgb = DANGER_RED

    r_cd1_sub = p_cd1.add_run()
    r_cd1_sub.text = "(Category: Fake KYC Phishing)"
    r_cd1_sub.font.name = FONT_MAIN
    r_cd1_sub.font.size = Pt(12)
    r_cd1_sub.font.bold = True
    r_cd1_sub.font.color.rgb = TEXT_TITLE

    p_cd2 = tf_cd.add_paragraph()
    p_cd2.space_before = Pt(4)
    r_cd2 = p_cd2.add_run()
    r_cd2.text = "INTERCEPTED TEXT: "
    r_cd2.font.name = FONT_MAIN
    r_cd2.font.size = Pt(10.5)
    r_cd2.font.bold = True
    r_cd2.font.color.rgb = TEXT_MUTED

    r_cd2_txt = p_cd2.add_run()
    r_cd2_txt.text = "\"Your bank account will be blocked today. Verify your KYC immediately using this link: http://sbi-kyc-update.net\""
    r_cd2_txt.font.name = FONT_MAIN
    r_cd2_txt.font.size = Pt(11)
    r_cd2_txt.font.italic = True
    r_cd2_txt.font.color.rgb = TEXT_BODY

    # Card 2: Why We Flagged This
    make_card(s7, Inches(3.8), Inches(3.45), Inches(8.733), Inches(1.65), CARD_BG, DANGER_RED, 1.5)
    tb_why = s7.shapes.add_textbox(Inches(4.0), Inches(3.52), Inches(8.3), Inches(1.5))
    tf_why = tb_why.text_frame
    tf_why.word_wrap = True

    p_wh = tf_why.paragraphs[0]
    r_wh = p_wh.add_run()
    r_wh.text = "🚩 WHY WE FLAGGED THIS (Explainable AI Indicators)"
    r_wh.font.name = FONT_MAIN
    r_wh.font.size = Pt(12)
    r_wh.font.bold = True
    r_wh.font.color.rgb = DANGER_RED

    flags = [
        "Threatens immediate account blockage or suspension to bypass standard caution",
        "Demands urgent KYC verification via an unverified external third-party domain",
        "Impersonates State Bank of India using deceptive URL 'sbi-kyc-update.net'",
        "High psychological pressure intended to force hasty action before verification"
    ]
    for f in flags:
        p_f = tf_why.add_paragraph()
        p_f.space_before = Pt(2)
        r_f = p_f.add_run()
        r_f.text = f"•  {f}"
        r_f.font.name = FONT_MAIN
        r_f.font.size = Pt(10.5)
        r_f.font.color.rgb = TEXT_BODY

    # Card 3: Actionable Guidance
    make_card(s7, Inches(3.8), Inches(5.25), Inches(8.733), Inches(1.65), CARD_BG, SAFE_GREEN, 1.5)
    tb_act = s7.shapes.add_textbox(Inches(4.0), Inches(5.32), Inches(8.3), Inches(1.5))
    tf_act = tb_act.text_frame
    tf_act.word_wrap = True

    p_ah = tf_act.paragraphs[0]
    r_ah = p_ah.add_run()
    r_ah.text = "✅ ACTIONABLE CITIZEN GUIDANCE (Zero-Jargon Decisions)"
    r_ah.font.name = FONT_MAIN
    r_ah.font.size = Pt(12)
    r_ah.font.bold = True
    r_ah.font.color.rgb = SAFE_GREEN

    actions = [
        ("WHAT TO DO: ", "Contact your bank branch using the official customer care number on your passbook/card.", SAFE_GREEN),
        ("WHAT TO DO: ", "Verify your account status strictly inside your official net banking app or branch.", SAFE_GREEN),
        ("WHAT NOT TO DO: ", "Do NOT click the link. Banks NEVER send external links for mandatory KYC.", DANGER_RED),
        ("WHAT NOT TO DO: ", "Do NOT share your OTP, UPI PIN, or passwords with anyone under any circumstance.", DANGER_RED)
    ]
    for prefix, act, col in actions:
        p_a = tf_act.add_paragraph()
        p_a.space_before = Pt(2)
        r_ap = p_a.add_run()
        r_ap.text = prefix
        r_ap.font.name = FONT_MAIN
        r_ap.font.size = Pt(10.5)
        r_ap.font.bold = True
        r_ap.font.color.rgb = col

        r_at = p_a.add_run()
        r_at.text = act
        r_at.font.name = FONT_MAIN
        r_at.font.size = Pt(10.5)
        r_at.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 8 — INNOVATION & USP (20 MARKS)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_bg(s8)
    add_header(s8, "07 | Innovation & Unique Value Proposition", "Why SCAMSHIELD Stands Apart.",
               "Judging Criterion: Innovation & Originality (20 Marks) — Redefining mobile threat prevention.")
    add_footer(s8, 8)

    make_card(s8, Inches(0.8), Inches(1.95), Inches(11.733), Inches(0.48), BLUE_BADGE, CYAN_PRIMARY, 1)
    tb_th = s8.shapes.add_textbox(Inches(1.0), Inches(1.98), Inches(11.333), Inches(0.4))
    tf_th = tb_th.text_frame
    p_th = tf_th.paragraphs[0]
    
    r_th1 = p_th.add_run()
    r_th1.text = "FEATURE / CAPABILITY"
    r_th1.font.name = FONT_MAIN
    r_th1.font.size = Pt(11)
    r_th1.font.bold = True
    r_th1.font.color.rgb = CYAN_PRIMARY

    r_th2 = p_th.add_run()
    r_th2.text = "                     TRADITIONAL CHECKERS"
    r_th2.font.name = FONT_MAIN
    r_th2.font.size = Pt(11)
    r_th2.font.bold = True
    r_th2.font.color.rgb = TEXT_MUTED

    r_th3 = p_th.add_run()
    r_th3.text = "                              SCAMSHIELD INNOVATION"
    r_th3.font.name = FONT_MAIN
    r_th3.font.size = Pt(11)
    r_th3.font.bold = True
    r_th3.font.color.rgb = SAFE_GREEN

    comp_rows = [
        ("Detection Mode", "Manual Copy-Paste (6+ friction steps)", "100% Zero-Touch Background Interception", SAFE_GREEN),
        ("Protection Timing", "Reactive (Only after user suspects something)", "Proactive (Instant heads-up alert before click)", SAFE_GREEN),
        ("Analysis Output", "Cryptic confidence scores ('84.2% phishing')", "Plain-Language Reasons + What To Do / Not To Do", SAFE_GREEN),
        ("User Privacy", "Cloud uploads of full messages / contacts", "Zero-Surveillance Edge Gate drops 90%+ benign chats", SAFE_GREEN),
        ("Accessibility", "English only, tiny text, complex dashboards", "High-contrast, Trilingual (EN/KN/HI), Audio TTS", SAFE_GREEN),
        ("Offline Reliability", "Crashes or fails completely without internet", "Dual-Engine: Cloud Groq AI + 100% Local Fallback", SAFE_GREEN)
    ]

    for i, (dim, trad, scam, col) in enumerate(comp_rows):
        top_r = Inches(2.55 + i * 0.6)
        make_card(s8, Inches(0.8), top_r, Inches(11.733), Inches(0.52), CARD_BG if i % 2 == 0 else CARD_BG_ALT, CARD_BORDER)
        
        tb_r = s8.shapes.add_textbox(Inches(1.0), top_r + Inches(0.08), Inches(11.333), Inches(0.38))
        tf_r = tb_r.text_frame
        p_r = tf_r.paragraphs[0]
        
        r_d = p_r.add_run()
        r_d.text = f"{dim:22} "
        r_d.font.name = FONT_MAIN
        r_d.font.size = Pt(11)
        r_d.font.bold = True
        r_d.font.color.rgb = TEXT_TITLE

        r_tr = p_r.add_run()
        r_tr.text = f"  {trad:45} "
        r_tr.font.name = FONT_MAIN
        r_tr.font.size = Pt(10.5)
        r_tr.font.color.rgb = TEXT_MUTED

        r_sc = p_r.add_run()
        r_sc.text = f"  {scam}"
        r_sc.font.name = FONT_MAIN
        r_sc.font.size = Pt(11)
        r_sc.font.bold = True
        r_sc.font.color.rgb = col

    make_card(s8, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.65), BLUE_BADGE, CYAN_PRIMARY, 1.5)
    tb_u = s8.shapes.add_textbox(Inches(1.0), Inches(6.3), Inches(11.333), Inches(0.55))
    tf_u = tb_u.text_frame
    p_u = tf_u.paragraphs[0]
    p_u.alignment = PP_ALIGN.CENTER
    r_u1 = p_u.add_run()
    r_u1.text = "UNIQUE VALUE PROPOSITION (USP):  "
    r_u1.font.name = FONT_MAIN
    r_u1.font.size = Pt(12)
    r_u1.font.bold = True
    r_u1.font.color.rgb = WARN_AMBER

    r_u2 = p_u.add_run()
    r_u2.text = "Zero-Touch Interception  +  Explainable Contextual AI  +  Zero-Surveillance Edge Gate  +  Elderly-First Accessibility."
    r_u2.font.name = FONT_MAIN
    r_u2.font.size = Pt(12)
    r_u2.font.bold = True
    r_u2.font.color.rgb = TEXT_TITLE

    # ==========================================
    # SLIDE 9 — TECHNOLOGY STACK
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_bg(s9)
    add_header(s9, "08 | Technology Stack", "Built on a Modern, Production-Grade Stack.",
               "Inspected and verified directly from the active SCAMSHIELD codebase.")
    add_footer(s9, 9)

    tech_stacks = [
        ("📱 ANDROID CLIENT", "Native Mobile Companion", CYAN_PRIMARY, [
            ("Language", "Kotlin 1.9 (Target SDK 34)"),
            ("UI Framework", "Jetpack Compose + Material 3"),
            ("Architecture", "Clean Architecture / MVVM + Coroutines & Flow"),
            ("Background Service", "Android NotificationListenerService & SmsReceiver"),
            ("Local Database", "Room SQLite Database (Encrypted metadata)"),
            ("Accessibility", "Android TextToSpeech (TTS) Engine"),
            ("Networking", "Retrofit 2 + OkHttp 3 with Dynamic Host")
        ]),
        ("⚡ BACKEND & AI INFERENCE", "Cloud Detection Engine", WARN_AMBER, [
            ("API Framework", "FastAPI (Python 3.14) + Uvicorn ASGI"),
            ("Cloud AI LLM", "Groq Cloud SDK (llama-3.3-70b-versatile)"),
            ("Inference Speed", "Sub-second contextual analysis (~350ms)"),
            ("Data Contract", "Pydantic v2 Type-Safe Request/Response Models"),
            ("Protocol", "RESTful JSON APIs with CORS Middleware"),
            ("Testing", "15 Comprehensive PyTest Test Cases"),
            ("Architecture", "Stateless, Highly Scalable Microservice")
        ]),
        ("🛡️ LOCAL EDGE ENGINE", "On-Device Privacy Gate", SAFE_GREEN, [
            ("Execution", "100% On-Device Regex & Weighted Heuristics"),
            ("Vectors Covered", "Indian Banking, UPI PIN, KYC, Lottery, Courier"),
            ("Latency", "< 5ms instant execution time"),
            ("False-Positive Shield", "Protective Pattern Checker ('Never share OTP' -> Safe)"),
            ("Zero-Surveillance", "Personal chats never leave the smartphone"),
            ("Deduplication", "Concurrent cache prevents repeat notifications"),
            ("Redundancy", "Guaranteed 100% offline fallback protection")
        ])
    ]

    for i, (title, sub, col, items) in enumerate(tech_stacks):
        left_c = Inches(0.8 + i * 3.96)
        make_card(s9, left_c, Inches(1.95), Inches(3.8), Inches(4.9), CARD_BG, col, 1.5)
        
        tb = s9.shapes.add_textbox(left_c + Inches(0.18), Inches(2.1), Inches(3.45), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_th = tf.paragraphs[0]
        r_th = p_th.add_run()
        r_th.text = title
        r_th.font.name = FONT_MAIN
        r_th.font.size = Pt(13)
        r_th.font.bold = True
        r_th.font.color.rgb = col

        p_ts = tf.add_paragraph()
        p_ts.space_before = Pt(2)
        r_ts = p_ts.add_run()
        r_ts.text = sub
        r_ts.font.name = FONT_MAIN
        r_ts.font.size = Pt(10)
        r_ts.font.color.rgb = TEXT_MUTED

        for label, val in items:
            p_i = tf.add_paragraph()
            p_i.space_before = Pt(7)
            r_il = p_i.add_run()
            r_il.text = f"{label}: "
            r_il.font.name = FONT_MAIN
            r_il.font.size = Pt(10.5)
            r_il.font.bold = True
            r_il.font.color.rgb = TEXT_TITLE

            r_iv = p_i.add_run()
            r_iv.text = val
            r_iv.font.name = FONT_MAIN
            r_iv.font.size = Pt(10)
            r_iv.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 10 — SYSTEM ARCHITECTURE & FEASIBILITY (25 MARKS)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_bg(s10)
    add_header(s10, "09 | System Architecture & Technical Feasibility", "Production Architecture Under the Hood.",
               "Judging Criterion: Technical Excellence & Implementation (25 Marks) — Robust, scalable & compliant.")
    add_footer(s10, 10)

    # Left: Editable System Architecture Flow Diagram with Real Shapes
    make_card(s10, Inches(0.8), Inches(1.90), Inches(6.1), Inches(5.0), CARD_BG, CYAN_PRIMARY, 1.5)
    
    tb_arch_head = s10.shapes.add_textbox(Inches(1.15), Inches(1.98), Inches(5.4), Inches(0.35))
    tf_ah = tb_arch_head.text_frame
    p_ah = tf_ah.paragraphs[0]
    r_ah = p_ah.add_run()
    r_ah.text = "SYSTEM ARCHITECTURE (EDITABLE PIPELINE)"
    r_ah.font.name = FONT_MAIN
    r_ah.font.size = Pt(11.5)
    r_ah.font.bold = True
    r_ah.font.color.rgb = CYAN_PRIMARY

    # Architecture flowchart nodes
    arch_nodes = [
        ("👤 USER / INCOMING CHATS", "SMS Broadcasts & WhatsApp Message Notifications", BLUE_BADGE, CYAN_PRIMARY),
        ("📱 MESSAGE DETECTION LAYER", "ScamNotificationListenerService + Android SmsReceiver", CARD_BG_ALT, CYAN_PRIMARY),
        ("⚡ LOCAL SCAM FILTER & DEDUP", "On-Device Regex (<5ms) + Shared Hash Deduplication", CARD_BG_ALT, SAFE_GREEN),
        ("☁️ BACKEND REST API", "FastAPI (Python 3.14) Microservice on Uvicorn ASGI", CARD_BG_ALT, WARN_AMBER),
        ("🧠 AI SCAM ANALYSIS & REASONING", "Groq Cloud SDK (LLaMA 3.3 70B) Intent & Phishing Scoring", CARD_BG_ALT, DANGER_RED),
        ("🛡️ CLASSIFICATION & UI DISPATCH", "0–100 Score + Heads-Up Alert + Trilingual UI + Voice TTS", BLUE_BADGE, SAFE_GREEN)
    ]

    for i, (n_title, n_sub, bg_c, border_c) in enumerate(arch_nodes):
        top_node = Inches(2.40 + i * 0.72)
        make_card(s10, Inches(1.0), top_node, Inches(5.7), Inches(0.56), bg_c, border_c, 1)
        
        tb_node = s10.shapes.add_textbox(Inches(1.15), top_node + Inches(0.04), Inches(5.4), Inches(0.48))
        tf_n = tb_node.text_frame
        p_nt = tf_n.paragraphs[0]
        
        r_nt = p_nt.add_run()
        r_nt.text = f"[{i+1}] {n_title}  "
        r_nt.font.name = FONT_MAIN
        r_nt.font.size = Pt(10)
        r_nt.font.bold = True
        r_nt.font.color.rgb = border_c

        p_ns = tf_n.add_paragraph()
        p_ns.space_before = Pt(1)
        r_ns = p_ns.add_run()
        r_ns.text = n_sub
        r_ns.font.name = FONT_MAIN
        r_ns.font.size = Pt(8.5)
        r_ns.font.color.rgb = TEXT_BODY

        # Add connecting downward arrow between nodes
        if i < len(arch_nodes) - 1:
            tb_arr = s10.shapes.add_textbox(Inches(3.6), top_node + Inches(0.53), Inches(0.5), Inches(0.2))
            tf_arr = tb_arr.text_frame
            p_arr = tf_arr.paragraphs[0]
            p_arr.alignment = PP_ALIGN.CENTER
            r_arr = p_arr.add_run()
            r_arr.text = "▼"
            r_arr.font.name = FONT_MAIN
            r_arr.font.size = Pt(7)
            r_arr.font.color.rgb = CYAN_PRIMARY

    # Right: Engineering Challenges Solved Card
    make_card(s10, Inches(7.15), Inches(1.95), Inches(5.383), Inches(4.9), CARD_BG, CARD_BORDER)
    tb_ch = s10.shapes.add_textbox(Inches(7.35), Inches(2.05), Inches(4.95), Inches(4.7))
    tf_ch = tb_ch.text_frame
    tf_ch.word_wrap = True

    p_chh = tf_ch.paragraphs[0]
    r_chh = p_chh.add_run()
    r_chh.text = "Key Technical Challenges & Solutions"
    r_chh.font.name = FONT_MAIN
    r_chh.font.size = Pt(14)
    r_chh.font.bold = True
    r_chh.font.color.rgb = WARN_AMBER

    challenges = [
        ("Background OS Lifecycle", "Android frequently unbinds listeners after updates. Solved with automatic rebind requests and PackageManager state cycling."),
        ("Cloud AI Latency & Cost", "Evaluating all chats in the cloud is slow and expensive. Dual-engine edge filter drops 90%+ benign messages on device instantly."),
        ("Duplicate Notifications", "WhatsApp frequently updates notifications. Implemented shared cross-channel hash deduplication across SMS and WhatsApp."),
        ("Offline Environments", "Users without active internet still need protection. Local heuristic engine provides 100% offline redundancy."),
        ("Preventing Infinite Loops", "SCAMSHIELD warning notifications are explicitly excluded from analysis to avoid recursion.")
    ]

    for c_title, c_sol in challenges:
        p_c = tf_ch.add_paragraph()
        p_c.space_before = Pt(8)
        r_ct = p_c.add_run()
        r_ct.text = f"⚡ {c_title}: "
        r_ct.font.name = FONT_MAIN
        r_ct.font.size = Pt(11)
        r_ct.font.bold = True
        r_ct.font.color.rgb = TEXT_TITLE

        r_cs = p_c.add_run()
        r_cs.text = c_sol
        r_cs.font.name = FONT_MAIN
        r_cs.font.size = Pt(10)
        r_cs.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 11 — REAL-WORLD IMPACT & FUTURE SCOPE (15 MARKS)
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_bg(s11)
    add_header(s11, "10 | Real-World Impact & Future Scope", "From One Smartphone to a Safer Digital Generation.",
               "Judging Criterion: Real-World Impact & Scalability (15 Marks) — Measurable protection & expansion roadmap.")
    add_footer(s11, 11)

    # Left: Current Real-World Impact
    make_card(s11, Inches(0.8), Inches(1.95), Inches(5.6), Inches(4.0), CARD_BG, SAFE_GREEN, 1.5)
    tb_imp = s11.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(5.1), Inches(3.8))
    tf_imp = tb_imp.text_frame
    tf_imp.word_wrap = True

    p_imh = tf_imp.paragraphs[0]
    r_imh = p_imh.add_run()
    r_imh.text = "IMMEDIATE REAL-WORLD IMPACT (LIVE TODAY)"
    r_imh.font.name = FONT_MAIN
    r_imh.font.size = Pt(12.5)
    r_imh.font.bold = True
    r_imh.font.color.rgb = SAFE_GREEN

    impacts = [
        ("Protecting Vulnerable Demographics", "Shields elderly parents and digitally inexperienced citizens from high-pressure deception."),
        ("Financial Loss Prevention", "Interprets fake KYC, UPI collect traps, and courier payment scams before funds are transferred."),
        ("Zero-Surveillance Privacy", "Maintains end-to-end user privacy by performing local edge filtering without storing personal messages."),
        ("Universal Voice Accessibility", "Text-to-Speech audio warnings bridge visual impairment and literacy barriers across regional demographics.")
    ]

    for title, desc in impacts:
        p_im = tf_imp.add_paragraph()
        p_im.space_before = Pt(8)
        r_it = p_im.add_run()
        r_it.text = f"✓ {title}: "
        r_it.font.name = FONT_MAIN
        r_it.font.size = Pt(10.5)
        r_it.font.bold = True
        r_it.font.color.rgb = TEXT_TITLE

        r_id = p_im.add_run()
        r_id.text = desc
        r_id.font.name = FONT_MAIN
        r_id.font.size = Pt(10)
        r_id.font.color.rgb = TEXT_BODY

    # Right: Future Scope
    make_card(s11, Inches(6.8), Inches(1.95), Inches(5.733), Inches(4.0), CARD_BG, CYAN_PRIMARY, 1.5)
    tb_fut = s11.shapes.add_textbox(Inches(7.05), Inches(2.05), Inches(5.2), Inches(3.8))
    tf_fut = tb_fut.text_frame
    tf_fut.word_wrap = True

    p_futh = tf_fut.paragraphs[0]
    r_futh = p_futh.add_run()
    r_futh.text = "FUTURE SCOPE & SCALE (ROADMAP)"
    r_futh.font.name = FONT_MAIN
    r_futh.font.size = Pt(12.5)
    r_futh.font.bold = True
    r_futh.font.color.rgb = CYAN_PRIMARY

    futures = [
        ("Broader Messaging Platforms", "Expanding listener modules to monitor Telegram, Signal, and Instagram Direct Messages."),
        ("Expanded Regional Languages", "Adding voice synthesis models for Bengali, Tamil, Telugu, and Marathi speakers."),
        ("Opt-in Family Guardian Network", "Alerting verified family guardians when a high-risk scam is detected on an elderly parent's device."),
        ("On-Device Small Language Models (SLMs)", "Deploying quantized local models via Google LiteRT/MediaPipe for 100% offline AI inference."),
        ("Real-Time Phishing Domain Sandboxing", "Deep inspection of shortened links with live domain reputation and phishing site classification.")
    ]

    for title, desc in futures:
        p_f = tf_fut.add_paragraph()
        p_f.space_before = Pt(6)
        r_ft = p_f.add_run()
        r_ft.text = f"○ {title}: "
        r_ft.font.name = FONT_MAIN
        r_ft.font.size = Pt(10.5)
        r_ft.font.bold = True
        r_ft.font.color.rgb = CYAN_PRIMARY

        r_fd = p_f.add_run()
        r_fd.text = desc
        r_fd.font.name = FONT_MAIN
        r_fd.font.size = Pt(9.5)
        r_fd.font.color.rgb = TEXT_BODY

    # Bottom: Expansion Scale Visual (4 Distinct Stage Cards)
    scale_steps = [
        ("1. SINGLE USER", "Real-Time Mobile Defense"),
        ("2. FAMILY CIRCLE", "Elderly Guardian Alerts"),
        ("3. COMMUNITY", "Regional Language Reach"),
        ("4. NATIONWIDE", "Scalable Proactive Defense")
    ]
    card_w = Inches(2.65)
    card_gap = Inches(0.37)
    start_x = Inches(0.8)
    for idx, (st_name, st_desc) in enumerate(scale_steps):
        cx = start_x + idx * (card_w + card_gap)
        make_card(s11, cx, Inches(6.12), card_w, Inches(0.76), BLUE_BADGE, CYAN_PRIMARY, 1)
        tb_sc = s11.shapes.add_textbox(cx + Inches(0.08), Inches(6.16), card_w - Inches(0.16), Inches(0.68))
        tf_sc = tb_sc.text_frame
        p_sc = tf_sc.paragraphs[0]
        p_sc.alignment = PP_ALIGN.CENTER
        r_sn = p_sc.add_run()
        r_sn.text = st_name + "\n"
        r_sn.font.name = FONT_MAIN
        r_sn.font.size = Pt(10)
        r_sn.font.bold = True
        r_sn.font.color.rgb = TEXT_TITLE

        r_sd = p_sc.add_run()
        r_sd.text = st_desc
        r_sd.font.name = FONT_MAIN
        r_sd.font.size = Pt(8.5)
        r_sd.font.color.rgb = CYAN_PRIMARY

        if idx < len(scale_steps) - 1:
            tb_arr = s11.shapes.add_textbox(cx + card_w, Inches(6.28), card_gap, Inches(0.4))
            tf_arr = tb_arr.text_frame
            p_arr = tf_arr.paragraphs[0]
            p_arr.alignment = PP_ALIGN.CENTER
            r_arr = p_arr.add_run()
            r_arr.text = "➔"
            r_arr.font.name = FONT_MAIN
            r_arr.font.size = Pt(12)
            r_arr.font.bold = True
            r_arr.font.color.rgb = WARN_AMBER

    # ==========================================
    # SLIDE 12 — FINAL PITCH / CONCLUSION
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    apply_bg(s12)

    tag12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.55), Inches(3.6), Inches(0.32))
    tag12.fill.solid()
    tag12.fill.fore_color.rgb = BLUE_BADGE
    tag12.line.color.rgb = CYAN_PRIMARY
    tag12.line.width = Pt(1)
    tf12 = tag12.text_frame
    p12 = tf12.paragraphs[0]
    p12.alignment = PP_ALIGN.CENTER
    r12 = p12.add_run()
    r12.text = "11 | CONCLUSION & FINAL PITCH"
    r12.font.name = FONT_MAIN
    r12.font.size = Pt(9.5)
    r12.font.bold = True
    r12.font.color.rgb = CYAN_PRIMARY

    tb_c12 = s12.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(8.0), Inches(0.48))
    tf_c12 = tb_c12.text_frame
    p_c12 = tf_c12.paragraphs[0]
    r_c12 = p_c12.add_run()
    r_c12.text = "Scams Shouldn't Be the User's Problem to Solve."
    r_c12.font.name = FONT_MAIN
    r_c12.font.size = Pt(23)
    r_c12.font.bold = True
    r_c12.font.color.rgb = TEXT_TITLE

    tb_c12_sub = s12.shapes.add_textbox(Inches(0.8), Inches(1.48), Inches(8.0), Inches(0.38))
    tf_c12_sub = tb_c12_sub.text_frame
    p_c12_sub = tf_c12_sub.paragraphs[0]
    r_c12_sub = p_c12_sub.add_run()
    r_c12_sub.text = "Building a digital ecosystem where suspicious messages trigger immediate protection — not panic."
    r_c12_sub.font.name = FONT_MAIN
    r_c12_sub.font.size = Pt(12)
    r_c12_sub.font.color.rgb = TEXT_MUTED

    concl_cards = [
        ("🛡️ 1. AUTOMATIC", "Zero user friction. Background detection intercepts threats before victims can be deceived."),
        ("🧠 2. EXPLAINABLE", "No cryptic scores. Plain-language indicators explain WHY a message is dangerous."),
        ("👴 3. ACCESSIBLE", "Built for elderly citizens first. Multilingual, high-contrast, and spoken voice guidance.")
    ]
    for i, (title, desc) in enumerate(concl_cards):
        left_cc = Inches(0.8 + i * 2.65)
        make_card(s12, left_cc, Inches(2.1), Inches(2.5), Inches(1.7), CARD_BG, CYAN_PRIMARY, 1)
        tb_cc = s12.shapes.add_textbox(left_cc + Inches(0.12), Inches(2.2), Inches(2.26), Inches(1.5))
        tf_cc = tb_cc.text_frame
        tf_cc.word_wrap = True
        p_cch = tf_cc.paragraphs[0]
        r_cch = p_cch.add_run()
        r_cch.text = title
        r_cch.font.name = FONT_MAIN
        r_cch.font.size = Pt(12.5)
        r_cch.font.bold = True
        r_cch.font.color.rgb = CYAN_PRIMARY

        p_ccd = tf_cc.add_paragraph()
        p_ccd.space_before = Pt(5)
        r_ccd = p_ccd.add_run()
        r_ccd.text = desc
        r_ccd.font.name = FONT_MAIN
        r_ccd.font.size = Pt(10)
        r_ccd.font.color.rgb = TEXT_BODY

    # Right Phone Mockup
    if os.path.exists(PHONE_DASHBOARD):
        s12.shapes.add_picture(PHONE_DASHBOARD, Inches(9.1), Inches(0.95), height=Inches(6.0))

    # Big Final Statement Card
    make_card(s12, Inches(0.8), Inches(4.05), Inches(7.8), Inches(1.5), BLUE_BADGE, CYAN_PRIMARY, 2)
    tb_bb = s12.shapes.add_textbox(Inches(1.0), Inches(4.18), Inches(7.4), Inches(1.2))
    tf_bb = tb_bb.text_frame
    tf_bb.word_wrap = True
    p_bb = tf_bb.paragraphs[0]
    p_bb.alignment = PP_ALIGN.CENTER
    r_bb1 = p_bb.add_run()
    r_bb1.text = "\"SCAMSHIELD DETECTS.\nSCAMSHIELD EXPLAINS.\nSCAMSHIELD PROTECTS.\""
    r_bb1.font.name = FONT_MAIN
    r_bb1.font.size = Pt(20)
    r_bb1.font.bold = True
    r_bb1.font.color.rgb = TEXT_TITLE

    # Team & Presentation Footer Card with Logo
    make_card(s12, Inches(0.8), Inches(5.8), Inches(7.8), Inches(1.05), CARD_BG, CARD_BORDER)
    
    if os.path.exists(LOGO_BADGE):
        s12.shapes.add_picture(LOGO_BADGE, Inches(0.95), Inches(5.9), width=Inches(0.85), height=Inches(0.85))

    tb_cr = s12.shapes.add_textbox(Inches(1.95), Inches(5.88), Inches(6.5), Inches(0.9))
    tf_cr = tb_cr.text_frame
    tf_cr.word_wrap = True
    p_cr = tf_cr.paragraphs[0]
    r_cr1 = p_cr.add_run()
    r_cr1.text = "SCAMSHIELD  •  Hackatopia 2K26 National Hackathon Finals\n"
    r_cr1.font.name = FONT_MAIN
    r_cr1.font.size = Pt(11)
    r_cr1.font.bold = True
    r_cr1.font.color.rgb = CYAN_PRIMARY

    p_cr2 = tf_cr.add_paragraph()
    p_cr2.space_before = Pt(3)
    r_cr2 = p_cr2.add_run()
    r_cr2.text = "Team: Aaron Kuriyan (Lead Developer)  •  @zer0neo (Collaborator)\n"
    r_cr2.font.name = FONT_MAIN
    r_cr2.font.size = Pt(10)
    r_cr2.font.color.rgb = TEXT_BODY

    r_cr3 = p_cr2.add_run()
    r_cr3.text = "Repository: Aaronkuriyan/scamshield-mobile  •  Ready for Live Demo"
    r_cr3.font.name = FONT_MAIN
    r_cr3.font.size = Pt(9.5)
    r_cr3.font.color.rgb = TEXT_MUTED

    # Save Presentation
    output_filename = "SCAMSHIELD_Hackatopia_Final_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully saved to: {os.path.abspath(output_filename)}")
    print(f"Total slides created: {len(prs.slides)}")

if __name__ == "__main__":
    build_presentation()
